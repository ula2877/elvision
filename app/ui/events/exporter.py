from __future__ import annotations

import csv
import time
import io
from pathlib import Path
from typing import Optional

from app.core.events.models import EventRecord
from app.core.logging import LogService


try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

try:
    from PIL import Image as PILImage
    HAS_PIL = True
except ImportError:
    HAS_PIL = False


_SNAPSHOT_COL_W = 13
_IMG_MAX_W = 64
_IMG_MAX_H = 48
_ROW_HEIGHT = 45

_COLUMNS = [
    ("ID", 8),
    ("Snapshot", _SNAPSHOT_COL_W),
    ("Timestamp", 22),
    ("Camera ID", 16),
    ("Camera Name", 20),
    ("Group", 16),
    ("Event Type", 18),
    ("Confidence", 12),
    ("Status", 14),
    ("Notes", 30),
    ("Snapshot Path", 40),
]

_COL_EVENT_TYPE = 6
_COL_STATUS = 8

_HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
_HEADER_ALIGN = Alignment(vertical="center", horizontal="center")
_HEADER_BORDER = Border(
    left=Side(style="thin", color="1F4E79"),
    right=Side(style="thin", color="1F4E79"),
    top=Side(style="thin", color="1F4E79"),
    bottom=Side(style="thin", color="1F4E79"),
)

_BODY_FONT = Font(color="000000", size=10)
_BODY_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
_BODY_ALIGN = Alignment(vertical="center", horizontal="left", wrap_text=False)
_BODY_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)

_STATUS_COLORS = {
    "new": Font(color="008000", size=10, bold=True),
    "acknowledged": Font(color="0070C0", size=10, bold=True),
    "resolved": Font(color="808080", size=10, bold=True),
    "dismissed": Font(color="FF0000", size=10, bold=True),
}

_EVENT_TYPE_COLORS = {
    "fall_detection": Font(color="FF0000", size=10, bold=True),
    "fire": Font(color="FF8C00", size=10, bold=True),
    "smoke": Font(color="800080", size=10, bold=True),
}


class EventExporter:
    def __init__(self) -> None:
        self._log = LogService.instance()

    def export_csv(self, events: list[EventRecord], path: str | Path) -> bool:
        try:
            path = Path(path)
            path.parent.mkdir(parents=True, exist_ok=True)
            with open(path, "w", newline="", encoding="utf-8-sig") as f:
                writer = csv.writer(f)
                writer.writerow([c[0] for c in _COLUMNS])
                for ev in events:
                    writer.writerow(self._row_data(ev))
            self._log.info("Exported %d events to CSV: %s", len(events), path)
            return True
        except Exception as exc:
            self._log.error("CSV export failed: %s", exc)
            return False

    def export_xlsx(self, events: list[EventRecord], path: str | Path) -> bool:
        if not HAS_OPENPYXL:
            self._log.error("openpyxl not available — install with: pip install openpyxl")
            return False

        try:
            path = Path(path)
            path.parent.mkdir(parents=True, exist_ok=True)
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Events"

            for col_idx, (_, width) in enumerate(_COLUMNS, 1):
                col_letter = get_column_letter(col_idx)
                ws.column_dimensions[col_letter].width = width

            for col_idx, (label, _) in enumerate(_COLUMNS, 1):
                cell = ws.cell(row=1, column=col_idx, value=label)
                cell.fill = _HEADER_FILL
                cell.font = _HEADER_FONT
                cell.alignment = _HEADER_ALIGN
                cell.border = _HEADER_BORDER

            ws.row_dimensions[1].height = 24
            snapshot_col = 2

            for row_idx, ev in enumerate(events, 2):
                data = self._row_data(ev)
                for col_idx, value in enumerate(data, 1):
                    cell = ws.cell(row=row_idx, column=col_idx, value=value)
                    cell.fill = _BODY_FILL
                    cell.alignment = _BODY_ALIGN
                    cell.border = _BODY_BORDER

                    if col_idx == _COL_EVENT_TYPE + 1:
                        cell.font = _EVENT_TYPE_COLORS.get(
                            ev.event_type, _BODY_FONT
                        )
                    elif col_idx == _COL_STATUS + 1:
                        cell.font = _STATUS_COLORS.get(
                            ev.status, _BODY_FONT
                        )
                    else:
                        cell.font = _BODY_FONT

                if ev.snapshot_path and HAS_PIL:
                    sp = Path(ev.snapshot_path)
                    if sp.exists():
                        try:
                            pil_img = PILImage.open(sp)
                            pil_img.thumbnail((_IMG_MAX_W, _IMG_MAX_H), PILImage.LANCZOS)
                            buffer = io.BytesIO()
                            pil_img.save(buffer, format="PNG")
                            buffer.seek(0)
                            img = XLImage(buffer)

                            ws.row_dimensions[row_idx].height = _ROW_HEIGHT

                            img_w_emu = int(img.width * 914400 / 96)
                            img_h_emu = int(img.height * 914400 / 96)

                            col_w_px = _SNAPSHOT_COL_W * 7
                            col_w_emu = int(col_w_px * 914400 / 96)
                            row_h_emu = int(_ROW_HEIGHT * 12700)

                            col_off = max(0, (col_w_emu - img_w_emu) // 2)
                            row_off = max(0, (row_h_emu - img_h_emu) // 2)

                            anchor = TwoCellAnchor()
                            anchor._from = AnchorMarker(
                                col=snapshot_col - 1,
                                row=row_idx - 1,
                                colOff=col_off,
                                rowOff=row_off,
                            )
                            anchor.to = AnchorMarker(
                                col=snapshot_col,
                                row=row_idx,
                                colOff=0,
                                rowOff=0,
                            )
                            img.anchor = anchor
                            ws.add_image(img)
                        except Exception:
                            pass

            # Auto-size text columns (skip Snapshot)
            for col_idx, (label, fixed_width) in enumerate(_COLUMNS, 1):
                if fixed_width == _SNAPSHOT_COL_W:
                    continue
                col_letter = get_column_letter(col_idx)
                max_len = len(label)
                for ev in events:
                    val = str(self._row_data(ev)[col_idx - 1])
                    max_len = max(max_len, len(val))
                ws.column_dimensions[col_letter].width = min(max_len + 3, 60)

            # Freeze header row
            ws.freeze_panes = "A2"

            # AutoFilter on header
            last_col = get_column_letter(len(_COLUMNS))
            ws.auto_filter.ref = f"A1:{last_col}1"

            # Convert to Excel Table
            data_last_row = len(events) + 1
            if data_last_row >= 1:
                table_ref = f"A1:{last_col}{data_last_row}"
                tab = Table(displayName="Events", ref=table_ref)
                style = TableStyleInfo(
                    name="TableStyleMedium9",
                    showFirstColumn=False,
                    showLastColumn=False,
                    showRowStripes=True,
                    showColumnStripes=False,
                )
                tab.tableStyleInfo = style
                ws.add_table(tab)

            wb.save(str(path))
            self._log.info("Exported %d events to XLSX: %s", len(events), path)
            return True
        except Exception as exc:
            self._log.error("XLSX export failed: %s", exc)
            return False

    @staticmethod
    def _row_data(ev: EventRecord) -> list:
        return [
            ev.id,
            "",
            time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(ev.timestamp)),
            ev.camera_id,
            ev.camera_name,
            ev.group_name,
            ev.event_label,
            f"{ev.confidence:.1%}" if ev.confidence > 0 else "",
            ev.status.capitalize(),
            ev.notes,
            ev.snapshot_path,
        ]
